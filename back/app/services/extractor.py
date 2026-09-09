from pathlib import Path

import pdfplumber
from pptx import Presentation


def extract_pdf(path: str) -> str:
    pages = []

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()

            if text:
                pages.append(text)

    return "\n\n".join(pages)


def extract_pptx(path: str) -> str:
    presentation = Presentation(path)

    slides = []

    for slide in presentation.slides:
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    texts.append(text)

        if texts:
            slides.append("\n".join(texts))

    return "\n\n".join(slides)


def extract_txt(path: str) -> str:
    return Path(path).read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_document(path: str) -> str:
    extension = Path(path).suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension in {".pptx", ".ppt"}:
        return extract_pptx(path)

    if extension == ".txt":
        return extract_txt(path)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )